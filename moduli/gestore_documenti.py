"""Lettura e ricerca locale nei documenti di Jarvis.

Non esegue i file, non li modifica e non invia contenuti a servizi remoti.
PDF/DOCX/PPTX richiedono i relativi pacchetti opzionali.
"""
from __future__ import annotations

import html
import os
import re
from pathlib import Path


ESTENSIONI_TESTO = {".txt", ".md", ".markdown", ".csv", ".json", ".jsonl", ".rst", ".log", ".py", ".html", ".htm"}
LIMITE_FILE_BYTE = 25 * 1024 * 1024
LIMITE_TESTO_CARATTERI = 50_000


class GestoreDocumentiJarvis:
    """Strumenti di sola lettura per documenti locali."""

    def estrai_testo(self, percorso: str) -> str:
        raw = str(percorso or "").strip().strip('"').strip("'").strip("“”")
        if not raw:
            raise ValueError("Indichi il percorso completo del file.")
        sorgente = Path(os.path.expanduser(raw)).resolve()
        if not sorgente.exists() or not sorgente.is_file():
            raise FileNotFoundError("Non trovo il file: " + str(sorgente))
        if sorgente.stat().st_size > LIMITE_FILE_BYTE:
            raise ValueError("Il file supera il limite di 25 MB per l'analisi locale.")

        estensione = sorgente.suffix.lower()
        if estensione in ESTENSIONI_TESTO:
            try:
                contenuto = sorgente.read_text(encoding="utf-8-sig")
            except UnicodeError as errore:
                raise ValueError("Il file testuale non è codificato in UTF-8.") from errore
            if estensione in {".html", ".htm"}:
                contenuto = re.sub(r"(?is)<(script|style)\b.*?</\1\s*>", " ", contenuto)
                contenuto = re.sub(r"(?s)<[^>]+>", " ", contenuto)
                contenuto = html.unescape(contenuto)
                contenuto = re.sub(r"[ \t]+", " ", contenuto)
            return contenuto[:LIMITE_TESTO_CARATTERI]

        if estensione == ".pdf":
            try:
                from pypdf import PdfReader
            except ImportError as errore:
                raise RuntimeError(
                    "Per leggere PDF installi il componente opzionale: "
                    "python -m pip install -r requirements-documenti.txt"
                ) from errore
            try:
                reader = PdfReader(str(sorgente))
                parti = []
                for pagina in reader.pages[:100]:
                    parti.append(pagina.extract_text() or "")
                return "\n".join(parti)[:LIMITE_TESTO_CARATTERI]
            except Exception as errore:
                raise ValueError("Non riesco a estrarre il testo dal PDF: " + str(errore)) from errore

        if estensione == ".docx":
            try:
                from docx import Document
            except ImportError as errore:
                raise RuntimeError(
                    "Per leggere DOCX installi il componente opzionale: "
                    "python -m pip install -r requirements-documenti.txt"
                ) from errore
            try:
                documento = Document(str(sorgente))
                parti = [p.text for p in documento.paragraphs if p.text.strip()]
                for tabella in documento.tables:
                    for riga in tabella.rows:
                        parti.append(" | ".join(cella.text.strip() for cella in riga.cells))
                return "\n".join(parti)[:LIMITE_TESTO_CARATTERI]
            except Exception as errore:
                raise ValueError("Non riesco a leggere il documento DOCX: " + str(errore)) from errore

        if estensione == ".pptx":
            try:
                from pptx import Presentation
            except ImportError as errore:
                raise RuntimeError(
                    "Per leggere PowerPoint installi: "
                    "python -m pip install -r requirements-presentazioni.txt"
                ) from errore
            try:
                presentazione = Presentation(str(sorgente))
                parti = []
                for slide in presentazione.slides:
                    for forma in slide.shapes:
                        if hasattr(forma, "text") and forma.text.strip():
                            parti.append(forma.text.strip())
                return "\n".join(parti)[:LIMITE_TESTO_CARATTERI]
            except Exception as errore:
                raise ValueError("Non riesco a leggere il PowerPoint: " + str(errore)) from errore

        raise ValueError(
            "Formato non supportato. Uso consentito: TXT, Markdown, CSV, JSON, "
            "JSONL, RST, LOG, PY, HTML, PDF, DOCX e PPTX."
        )

    def leggi_file(self, percorso: str, limite_caratteri: int = 5000) -> str:
        """Restituisce una porzione leggibile senza eseguire o modificare il file."""
        try:
            contenuto = self.estrai_testo(percorso).strip()
        except (ValueError, RuntimeError, FileNotFoundError, OSError) as errore:
            return str(errore)
        if not contenuto:
            return "Il file non contiene testo estraibile."
        limite = max(500, min(int(limite_caratteri), 8000))
        troncato = len(contenuto) > limite
        risultato = contenuto[:limite].rstrip()
        if troncato:
            risultato += "\n\n[Contenuto abbreviato. Per vedere il resto, consulti il file originale.]"
        return risultato

    def cerca_nel_file(self, percorso: str, termine: str) -> str:
        """Cerca testo letterale, senza interpretarlo come espressione o comando."""
        termine = str(termine or "").strip()
        if not termine:
            return "Indichi il testo da cercare."
        if len(termine) > 200:
            return "Il testo da cercare è troppo lungo (massimo 200 caratteri)."
        try:
            contenuto = self.estrai_testo(percorso)
        except (ValueError, RuntimeError, FileNotFoundError, OSError) as errore:
            return str(errore)
        righe = contenuto.splitlines()
        trovati = []
        for numero, riga in enumerate(righe, 1):
            posizione = riga.casefold().find(termine.casefold())
            if posizione >= 0:
                inizio = max(0, posizione - 70)
                fine = min(len(riga), posizione + len(termine) + 100)
                frammento = riga[inizio:fine].strip()
                if inizio:
                    frammento = "…" + frammento
                if fine < len(riga):
                    frammento += "…"
                trovati.append(f"Riga {numero}: {frammento}")
                if len(trovati) >= 10:
                    break
        if not trovati:
            return f"Non ho trovato «{termine}» nel file."
        intestazione = f"Ho trovato «{termine}» in {Path(os.path.expanduser(percorso.strip().strip('"').strip("'"))).name}:"
        return intestazione + "\n" + "\n".join(trovati)
