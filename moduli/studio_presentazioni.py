"""Studio presentazioni locale e opzionale per Jarvis.

Crea file PowerPoint modificabili a partire da testo o Markdown fornito
dall'utente. Non usa servizi cloud e non sovrascrive presentazioni esistenti.
Richiede il pacchetto opzionale python-pptx.
"""
from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Iterable


BASE_JARVIS = Path(__file__).resolve().parent.parent
CARTELLA_PRESENTAZIONI = BASE_JARVIS / "outputs" / "presentazioni"


class StudioPresentazioniJarvis:
    """Genera presentazioni .pptx da contenuti locali, senza rete."""

    def __init__(self, cartella_output: str | os.PathLike | None = None):
        self.cartella_output = Path(cartella_output) if cartella_output else CARTELLA_PRESENTAZIONI

    @staticmethod
    def analizza_testo(contenuto: str, titolo_predefinito: str = "Presentazione") -> list[dict[str, object]]:
        """Converte testo semplice/Markdown in diapositive strutturate."""
        testo = str(contenuto or "").replace("\r\n", "\n").replace("\r", "\n").strip()
        if not testo:
            return []

        diapositive: list[dict[str, object]] = []
        titolo = titolo_predefinito.strip() or "Presentazione"
        corpo: list[str] = []

        def aggiungi():
            nonlocal corpo
            righe = [riga.strip() for riga in corpo if riga.strip()]
            if righe:
                diapositive.append({"titolo": titolo, "punti": righe[:8]})
            corpo = []

        for riga in testo.split("\n"):
            pulita = riga.strip()
            intestazione = re.match(r"^#{1,3}\s+(.+)$", pulita)
            if intestazione:
                aggiungi()
                titolo = intestazione.group(1).strip()[:120] or titolo_predefinito
                continue
            if not pulita:
                continue
            pulita = re.sub(r"^(?:[-*+]\s+|\d+[.)]\s+)", "", pulita).strip()
            # Un paragrafo lungo viene spezzato in punti leggibili.
            if len(pulita) > 220:
                parole = pulita.split()
                segmento = ""
                for parola in parole:
                    if segmento and len(segmento) + len(parola) + 1 > 180:
                        corpo.append(segmento)
                        segmento = parola
                    else:
                        segmento = (segmento + " " + parola).strip()
                if segmento:
                    corpo.append(segmento)
            else:
                corpo.append(pulita)
            if len(corpo) >= 8:
                aggiungi()
        aggiungi()

        if not diapositive:
            return []
        return diapositive

    def crea_da_file(self, percorso: str) -> str:
        """Legge solo file testuali supportati e crea una copia .pptx."""
        raw = str(percorso or "").strip().strip('"').strip("'")
        if not raw:
            return "Indichi il percorso completo di un file .txt o .md."
        sorgente = Path(os.path.expanduser(raw)).resolve()
        if not sorgente.is_file():
            return f"Non trovo il file: {sorgente}"
        if sorgente.suffix.lower() not in {".txt", ".md", ".markdown"}:
            return "Per ora posso creare presentazioni da file .txt, .md o .markdown."
        try:
            contenuto = sorgente.read_text(encoding="utf-8")
        except (OSError, UnicodeError) as errore:
            return f"Non riesco a leggere il file in UTF-8: {errore}"
        return self.crea_da_contenuto(contenuto, sorgente.stem)

    def crea_da_contenuto(self, contenuto: str, titolo: str = "Presentazione") -> str:
        """Crea una presentazione modificabile; l'output non viene mai sovrascritto."""
        diapositive = self.analizza_testo(contenuto, titolo)
        if not diapositive:
            return "Il contenuto è vuoto: aggiunga testo o titoli Markdown."
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.enum.text import PP_ALIGN
        except ImportError:
            return (
                "Per creare PowerPoint serve il componente opzionale. "
                "Dalla cartella Jarvis esegua: python -m pip install -r requirements-presentazioni.txt"
            )

        try:
            self.cartella_output.mkdir(parents=True, exist_ok=True)
            presentazione = Presentation()
            presentazione.slide_width = Inches(13.333)
            presentazione.slide_height = Inches(7.5)

            prima = presentazione.slides.add_slide(presentazione.slide_layouts[0])
            prima.shapes.title.text = self._titolo_pulito(titolo)
            sottotitolo = prima.placeholders[1]
            sottotitolo.text = "Preparata localmente da Jarvis"
            sottotitolo.text_frame.paragraphs[0].font.size = Pt(18)

            for elemento in diapositive:
                slide = presentazione.slides.add_slide(presentazione.slide_layouts[1])
                slide.shapes.title.text = str(elemento["titolo"])[:120]
                frame = slide.placeholders[1].text_frame
                frame.clear()
                for indice, punto in enumerate(elemento["punti"]):
                    paragrafo = frame.paragraphs[0] if indice == 0 else frame.add_paragraph()
                    paragrafo.text = str(punto)[:500]
                    paragrafo.level = 0
                    paragrafo.font.size = Pt(22 if len(str(punto)) < 100 else 18)

            nome_base = re.sub(r"[^A-Za-z0-9._-]+", "_", self._titolo_pulito(titolo)).strip("._-") or "Presentazione"
            destinazione = self.cartella_output / (nome_base[:70] + ".pptx")
            numero = 2
            while destinazione.exists():
                destinazione = self.cartella_output / (nome_base[:65] + f"_{numero}.pptx")
                numero += 1
            presentazione.save(str(destinazione))
            return f"Presentazione PowerPoint creata con {len(diapositive) + 1} diapositive: {destinazione}"
        except Exception as errore:
            return f"Non sono riuscito a creare la presentazione: {errore}"

    @staticmethod
    def _titolo_pulito(titolo: str) -> str:
        return " ".join(str(titolo or "Presentazione").split())[:120] or "Presentazione"


__all__ = ["StudioPresentazioniJarvis"]
